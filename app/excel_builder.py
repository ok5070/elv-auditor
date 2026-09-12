import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import Dict, Any

def generate_bom_excel(project_name: str, client_name: str, bom: Dict[str, Any]) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Спецификация BOM"

    ws.views.sheetView[0].showGridLines = True

    title_font = Font(name="Calibri", size=14, bold=True, color="1F497D")
    meta_font = Font(name="Calibri", size=11, italic=True)
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    alt_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    ws["A1"] = f"Спецификация оборудования и материалов: {project_name}"
    ws["A1"].font = title_font
    ws["A2"] = f"Заказчик: {client_name}"
    ws["A2"].font = meta_font

    headers = ["№", "Наименование позиции", "Количество", "Ед. изм."]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[4].height = 24

    current_row = 5
    for idx, (item_name, quantity) in enumerate(bom.items(), 1):
        unit = "шт"
        name_lower = item_name.lower()
        if "бухт" in name_lower:
            unit = "бухта"
        elif "метр" in name_lower or "гофра" in name_lower or "кабель" in name_lower:
            unit = "м"

        c_idx = ws.cell(row=current_row, column=1, value=idx)
        c_name = ws.cell(row=current_row, column=2, value=item_name)
        c_qty = ws.cell(row=current_row, column=3, value=quantity)
        c_unit = ws.cell(row=current_row, column=4, value=unit)

        c_idx.alignment = Alignment(horizontal="center")
        c_name.alignment = Alignment(horizontal="left")
        c_qty.alignment = Alignment(horizontal="right")
        c_unit.alignment = Alignment(horizontal="center")

        for c in (c_idx, c_name, c_qty, c_unit):
            c.font = data_font
            c.border = thin_border
            if current_row % 2 == 0:
                c.fill = alt_fill

        ws.row_dimensions[current_row].height = 20
        current_row += 1

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 10)

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 50

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
