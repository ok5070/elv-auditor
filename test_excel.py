import urllib.request
import json
import os
import openpyxl

payload = {
    "project_name": "Офис БЦ Сенатор",
    "client_name": "ООО Сибест",
    "rack_pos": {"x": 10.0, "y": 20.0, "z": 2.2},
    "ceiling_tray_height": 3.5,
    "doors": [
        {
            "door_id": "D-01",
            "name": "Главный эвакуационный выход",
            "width_mm": 950,
            "is_evacuation": True,
            "from_zone": "Коридор",
            "to_zone": "Улица"
        }
    ],
    "cameras": [
        {
            "camera_id": "CAM-01",
            "name": "Входная группа",
            "pos": {"x": 1.0, "y": 1.0, "z": 3.0},
            "dist_to_wall": 2.5,
            "focal_length_mm": 2.8,
            "resolution_w": 1920
        }
    ],
    "wifi_points": [
        {
            "ap_id": "AP-01",
            "pos": {"x": 5.0, "y": 5.0},
            "target_zone_name": "OpenSpace",
            "target_zone_clients": 20,
            "test_distance_m": 10.0
        }
    ],
    "audio_zones": [
        {
            "zone_id": "AZ-01",
            "name": "Фоновый звук Офис",
            "area_sqm": 80.0,
            "ceiling_height_m": 3.2,
            "system_type": "100V",
            "cable_distance_to_rack_m": 40.0
        }
    ]
}

req = urllib.request.Request(
    "http://127.0.0.1:8085/api/v1/audit/excel",
    data=json.dumps(payload).encode("utf-8"),
    headers={"Content-Type": "application/json"}
)

output_file = "BOM_Office_Senator.xlsx"

try:
    with urllib.request.urlopen(req) as resp:
        with open(output_file, "wb") as f:
            f.write(resp.read())
    
    print(f"1. Файл {output_file} получен из API ({os.path.getsize(output_file)} байт).")
    
    wb = openpyxl.load_workbook(output_file)
    ws = wb.active
    print(f"2. Имя листа: {ws.title}")
    print(f"3. {ws['A1'].value}")
    print(f"4. {ws['A2'].value}")
    print("\n5. Содержимое сформированной спецификации:")
    for r in range(4, ws.max_row + 1):
        row_vals = [str(ws.cell(row=r, column=c).value) for c in range(1, 5)]
        print("   | " + " | ".join(row_vals) + " |")
    print("========================================")
except Exception as e:
    print("Ошибка получения Excel:", e)
