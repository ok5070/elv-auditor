import urllib.request
import os
import openpyxl
import ezdxf

dxf_file = "pipeline_test.dxf"
doc = ezdxf.new(dxfversion="R2010")
msp = doc.modelspace()

# Комната 1 (Офис 10x8 = 80 кв.м)
msp.add_lwpolyline([(0, 0), (10, 0), (10, 8), (0, 8)], close=True, dxfattribs={"layer": "A-WALL"})
# Комната 2 (Конференц-зал 6x6 = 36 кв.м)
msp.add_lwpolyline([(12, 0), (18, 0), (18, 6), (12, 6)], close=True, dxfattribs={"layer": "A-WALL"})
doc.saveas(dxf_file)

# Отправляем multipart/form-data
boundary = "----WebKitFormBoundaryELVAuditorPipeline"
body = []

body.append(f"--{boundary}".encode("utf-8"))
body.append(b'Content-Disposition: form-data; name="project_name"\r\n')
body.append("БЦ Аврора - Автоматический расчет по DXF".encode("utf-8"))

body.append(f"--{boundary}".encode("utf-8"))
body.append(b'Content-Disposition: form-data; name="client_name"\r\n')
body.append("ООО Сибест".encode("utf-8"))

body.append(f"--{boundary}".encode("utf-8"))
body.append(f'Content-Disposition: form-data; name="file"; filename="{dxf_file}"\r\nContent-Type: application/dxf\r\n'.encode("utf-8"))
with open(dxf_file, "rb") as f:
    body.append(f.read())

body.append(f"--{boundary}--\r\n".encode("utf-8"))
payload_bytes = b"\r\n".join(body)

req = urllib.request.Request(
    "http://127.0.0.1:8085/api/v1/dxf/audit-excel",
    data=payload_bytes,
    headers={"Content-Type": f"multipart/form-data; boundary={boundary}"}
)

output_file = "BOM_Auto_Pipeline.xlsx"
try:
    with urllib.request.urlopen(req) as resp:
        with open(output_file, "wb") as f:
            f.write(resp.read())
    
    print(f"1. Спецификация успешно создана из DXF: {output_file} ({os.path.getsize(output_file)} байт).")
    
    wb = openpyxl.load_workbook(output_file)
    ws = wb.active
    print(f"2. {ws['A1'].value}")
    print(f"3. {ws['A2'].value}")
    print("\n4. Итоговая смета по чертежу:")
    for r in range(4, ws.max_row + 1):
        row_vals = [str(ws.cell(row=r, column=c).value) for c in range(1, 5)]
        print("   | " + " | ".join(row_vals) + " |")
    print("========================================")
except Exception as e:
    print("Ошибка пайплайна:", e)
